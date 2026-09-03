"""Utility functions for the application."""
import re
import zipfile
from pathlib import Path
from typing import Optional

from fastapi import HTTPException


# ============================================================
# Skill Normalization
# ============================================================

SKILL_ALIASES = {
    # Programming Languages
    "python": "Python",
    "py": "Python",
    "java": "Java",
    "javascript": "JavaScript",
    "js": "JavaScript",
    "typescript": "TypeScript",
    "ts": "TypeScript",
    "c++": "C++",
    "cpp": "C++",
    "c#": "C#",
    "csharp": "C#",
    "ruby": "Ruby",
    "go": "Go",
    "golang": "Go",
    "rust": "Rust",
    "php": "PHP",
    "html": "HTML",
    "css": "CSS",
    "jquery": "jQuery",
    "bootstrap": "Bootstrap",
    "tailwind": "Tailwind CSS",
    "tailwind css": "Tailwind CSS",
    "c/c++": "C++",
    "data structures": "Data Structures",
    "figma": "Figma",
    
    # Web Frameworks
    "react": "React",
    "reactjs": "React",
    "react.js": "React",
    "vue": "Vue",
    "vuejs": "Vue",
    "vue.js": "Vue",
    "angular": "Angular",
    "node": "Node.js",
    "nodejs": "Node.js",
    "node.js": "Node.js",
    "django": "Django",
    "flask": "Flask",
    "fastapi": "FastAPI",
    "spring": "Spring Boot",
    "springboot": "Spring Boot",
    
    # Data & ML
    "machine learning": "Machine Learning",
    "ml": "Machine Learning",
    "deep learning": "Deep Learning",
    "ai": "Artificial Intelligence",
    "artificial intelligence": "Artificial Intelligence",
    "tensorflow": "TensorFlow",
    "pytorch": "PyTorch",
    "keras": "Keras",
    "sklearn": "Scikit-learn",
    "scikit-learn": "Scikit-learn",
    "scikit learn": "Scikit-learn",
    "pandas": "Pandas",
    "numpy": "NumPy",
    "scipy": "SciPy",
    "nlp": "NLP",
    "natural language processing": "NLP",
    "computer vision": "Computer Vision",
    "cv": "Computer Vision",
    
    # Databases
    "sql": "SQL",
    "mysql": "MySQL",
    "postgres": "PostgreSQL",
    "postgresql": "PostgreSQL",
    "mongodb": "MongoDB",
    "redis": "Redis",
    "cassandra": "Cassandra",
    "elasticsearch": "Elasticsearch",
    "dynamodb": "DynamoDB",
    
    # Cloud & DevOps
    "aws": "AWS",
    "amazon web services": "AWS",
    "azure": "Azure",
    "gcp": "Google Cloud Platform",
    "google cloud": "Google Cloud Platform",
    "docker": "Docker",
    "kubernetes": "Kubernetes",
    "k8s": "Kubernetes",
    "jenkins": "Jenkins",
    "git": "Git",
    "github": "GitHub",
    "gitlab": "GitLab",
    "ci/cd": "CI/CD",
    "devops": "DevOps",
    
    # Tools & Others
    "linux": "Linux",
    "unix": "Unix",
    "windows": "Windows",
    "macos": "macOS",
    "docker": "Docker",
    "rest": "REST APIs",
    "rest api": "REST APIs",
    "graphql": "GraphQL",
    "statistics": "Statistics",
    "statistics": "Statistics",
}

SKILL_CATEGORIES = {
    "Python": "Programming Languages",
    "Java": "Programming Languages",
    "JavaScript": "Programming Languages",
    "TypeScript": "Programming Languages",
    "C++": "Programming Languages",
    "C#": "Programming Languages",
    "Ruby": "Programming Languages",
    "Go": "Programming Languages",
    "Rust": "Programming Languages",
    "PHP": "Programming Languages",
    "HTML": "Web Technologies",
    "CSS": "Web Technologies",
    "jQuery": "Web Frameworks",
    "Bootstrap": "Web Frameworks",
    "Tailwind CSS": "Web Frameworks",
    "Data Structures": "Computer Science",
    "Figma": "Tools & Design",
    
    "React": "Web Frameworks",
    "Vue": "Web Frameworks",
    "Angular": "Web Frameworks",
    "Node.js": "Web Frameworks",
    "Django": "Web Frameworks",
    "Flask": "Web Frameworks",
    "FastAPI": "Web Frameworks",
    "Spring Boot": "Web Frameworks",
    
    "Machine Learning": "ML & AI",
    "Deep Learning": "ML & AI",
    "Artificial Intelligence": "ML & AI",
    "TensorFlow": "ML & AI",
    "PyTorch": "ML & AI",
    "Keras": "ML & AI",
    "Scikit-learn": "ML & AI",
    "NLP": "ML & AI",
    "Computer Vision": "ML & AI",
    
    "SQL": "Database",
    "MySQL": "Database",
    "PostgreSQL": "Database",
    "MongoDB": "Database",
    "Redis": "Database",
    "Cassandra": "Database",
    "Elasticsearch": "Database",
    "DynamoDB": "Database",
    
    "AWS": "Cloud & DevOps",
    "Azure": "Cloud & DevOps",
    "Google Cloud Platform": "Cloud & DevOps",
    "Docker": "Cloud & DevOps",
    "Kubernetes": "Cloud & DevOps",
    "Jenkins": "Cloud & DevOps",
    "CI/CD": "Cloud & DevOps",
    "DevOps": "Cloud & DevOps",
    
    "Git": "Tools & Version Control",
    "GitHub": "Tools & Version Control",
    "GitLab": "Tools & Version Control",
    "Linux": "Operating Systems",
    "Unix": "Operating Systems",
    "Windows": "Operating Systems",
    "macOS": "Operating Systems",
    "REST APIs": "API & Communication",
    "GraphQL": "API & Communication",
    "Statistics": "Data Science",
    "Pandas": "Data Science",
    "NumPy": "Data Science",
}


def normalize_skill(skill_raw: str) -> str:
    """Normalize skill name using aliases."""
    normalized = SKILL_ALIASES.get(skill_raw.lower().strip())
    if normalized:
        return normalized
    return skill_raw.strip().title()


def get_skill_category(skill: str) -> str:
    """Get category for a skill."""
    return SKILL_CATEGORIES.get(skill, "Technical")


def extract_skills_from_text(text: str) -> list[tuple[str, str]]:
    """
    Extract skills from text.
    
    Returns list of (skill_name, category) tuples.
    """
    # Use the full text (with newlines) for better section-aware matching
    text_lower = text.lower()
    found_skills = []
    found_skill_names = set()
    
    # Sort aliases by length descending to match longer phrases first (e.g. "machine learning" before "ml")
    sorted_aliases = sorted(SKILL_ALIASES.items(), key=lambda x: len(x[0]), reverse=True)
    
    # Search for each skill alias in the text
    for alias, normalized_name in sorted_aliases:
        # Use non-word boundaries so names containing symbols (for example C/C++) match.
        escaped = re.escape(alias)
        pattern = r"(?i)(?<!\w)" + escaped + r"(?!\w)"
        if re.search(pattern, text_lower):
            if normalized_name not in found_skill_names:
                category = get_skill_category(normalized_name)
                found_skills.append((normalized_name, category))
                found_skill_names.add(normalized_name)
    
    return found_skills


# ============================================================
# Section Detection
# ============================================================

SECTION_KEYWORDS = {
    "education": [
        "education", "education and training", "academic background",
        "degree", "university", "college", "school", "certification",
        "certifications", "courses", "training"
    ],
    "experience": [
        "experience", "work experience", "professional experience",
        "employment", "employment history", "career history",
        "positions held", "job history", "work history"
    ],
    "projects": [
        "projects", "portfolio", "projects and portfolio",
        "notable projects", "key projects", "professional projects",
        "personal projects", "opensource"
    ],
    "skills": [
        "skills", "technical skills", "core competencies",
        "competencies", "expertise", "technologies"
    ],
    "certifications": [
        "certifications", "certified", "certification",
        "credentials", "licenses", "professional certifications"
    ],
    "achievements": [
        "achievements", "accomplishments", "awards",
        "honors", "awards and honors"
    ],
}


def detect_sections(text: str) -> dict[str, list[str]]:
    """Detect resume sections and extract content."""
    lines = text.split("\n")
    sections = {
        "education": [],
        "experience": [],
        "projects": [],
        "skills": [],
        "certifications": [],
        "achievements": [],
    }
    
    current_section = None
    
    for line in lines:
        line_clean = line.strip()
        line_lower = line_clean.lower()
        
        # Check if this line is a section header
        found_section = False
        for section_name, keywords in SECTION_KEYWORDS.items():
            for keyword in keywords:
                if keyword in line_lower:
                    current_section = section_name
                    found_section = True
                    break
            if found_section:
                break
        
        # Add content to current section
        if current_section and line_clean and not found_section:
            sections[current_section].append(line_clean)
    
    return sections


# ============================================================
# File Handling
# ============================================================

def safe_filename(filename: str) -> str:
    """Generate safe filename from user input."""
    # Remove path separators and special characters
    filename = re.sub(r"[^\w\s\.\-]", "", filename)
    # Replace multiple dots with single dot
    filename = re.sub(r"\.{2,}", ".", filename)
    # Limit length
    filename = filename[:255]
    return filename or "resume"


def validate_file_extension(filepath: Path, allowed_extensions: tuple) -> bool:
    """Validate file extension."""
    return filepath.suffix.lower() in allowed_extensions


def validate_file_size(
    file_size: int,
    max_size_bytes: int
) -> bool:
    """Validate file size."""
    return 0 < file_size <= max_size_bytes


# ============================================================
# Resume Parsing
# ============================================================

def parse_pdf(filepath: Path) -> str:
    """Extract text from PDF file, preserving line structure."""
    try:
        import fitz  # PyMuPDF
        
        pdf = fitz.open(filepath)
        text_parts = []
        
        for page_num in range(len(pdf)):
            page = pdf[page_num]
            # Use "text" mode which preserves word boundaries and newlines
            text = page.get_text("text")
            if text.strip():
                text_parts.append(text)
        
        full_text = "\n".join(text_parts)
        # If page.get_text gave nothing (image PDF), return empty string cleanly
        return full_text if full_text.strip() else ""
    except Exception as error:
        raise HTTPException(
            status_code=422,
            detail={
                "code": "RESUME_PARSE_ERROR",
                "message": f"Unable to parse PDF: {str(error)}"
            }
        )


def parse_docx(filepath: Path) -> str:
    """Extract text from DOCX file."""
    try:
        if not zipfile.is_zipfile(filepath):
            raw = filepath.read_bytes()
            if raw:
                try:
                    return raw.decode("utf-8")
                except UnicodeDecodeError:
                    return raw.decode("latin-1", errors="ignore")

        from docx import Document

        doc = Document(filepath)
        text_parts = []

        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)

        # Also process tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text)

        return "\n".join(text_parts)
    except Exception as error:
        raise HTTPException(
            status_code=422,
            detail={
                "code": "RESUME_PARSE_ERROR",
                "message": f"Unable to parse DOCX: {str(error)}"
            }
        )


def parse_document(filepath: Path, file_type: str) -> str:
    """Parse resume document and extract text."""
    if not filepath.exists():
        raise HTTPException(
            status_code=404,
            detail={
                "code": "FILE_NOT_FOUND",
                "message": "Resume file not found"
            }
        )
    
    suffix = filepath.suffix.lower()
    
    if suffix == ".pdf" or "pdf" in file_type.lower():
        return parse_pdf(filepath)
    elif suffix in (".docx", ".doc") or "wordprocessing" in file_type.lower():
        return parse_docx(filepath)
    else:
        # Try to read as plain text
        try:
            return filepath.read_text(encoding="utf-8", errors="ignore")
        except Exception as error:
            raise HTTPException(
                status_code=422,
                detail={
                    "code": "RESUME_PARSE_ERROR",
                    "message": f"Unsupported file format: {file_type}"
                }
            )


# ============================================================
# Text Cleaning
# ============================================================

def clean_text(text: str) -> str:
    """Clean and normalize resume text while preserving line structure."""
    # Remove control characters but KEEP newlines (\n) and tabs (\t) for section detection
    text = re.sub(r"[\x00-\x08\x0b-\x0c\x0e-\x1f\x7f-\x9f]", "", text)
    # Collapse multiple spaces on same line to single space (but don't touch newlines)
    text = re.sub(r"[ \t]+", " ", text)
    # Collapse 3+ consecutive newlines to 2
    text = re.sub(r"\n{3,}", "\n\n", text)
    # Normalize quotes
    text = text.replace("\u201c", '"').replace("\u201d", '"').replace("\u2019", "'")
    
    return text.strip()


# ============================================================
# Score Calculation
# ============================================================

def calculate_resume_score_breakdown(
    text: str,
    skills_found: int,
    sections: dict[str, list[str]]
) -> dict[str, int]:
    """Calculate detailed resume score breakdown."""
    # Skills score: based on number of skills (0-100)
    skills_score = min(100, max(0, skills_found * 12))
    
    # Experience score: based on experience section
    experience_score = min(100, len(sections.get("experience", [])) * 3 + 40)
    
    # Education score: based on education section
    education_score = min(100, len(sections.get("education", [])) * 10 + 30)
    
    # Projects score: based on projects section
    projects_score = min(100, len(sections.get("projects", [])) * 15 + 20)
    
    # Keywords score: based on technical terms
    keywords_count = len(
        re.findall(
            r"\b(api|database|cloud|automation|algorithm|architecture|deployment|testing|optimization)\b",
            text.lower()
        )
    )
    keywords_score = min(100, keywords_count * 10)
    
    return {
        "skills": skills_score,
        "experience": experience_score,
        "education": education_score,
        "projects": projects_score,
        "keywords": keywords_score,
    }


def calculate_overall_resume_score(breakdown: dict[str, int]) -> int:
    """Calculate overall resume score from breakdown."""
    # Weighted average: Skills 30%, Experience 25%, Education 15%, Projects 20%, Keywords 10%
    weights = {
        "skills": 0.30,
        "experience": 0.25,
        "education": 0.15,
        "projects": 0.20,
        "keywords": 0.10,
    }
    
    overall = sum(breakdown.get(key, 0) * weight for key, weight in weights.items())
    return int(round(overall))


# ============================================================
# Validation
# ============================================================

def get_error_response(code: str, message: str) -> dict:
    """Create standard error response."""
    return {
        "success": False,
        "error": {
            "code": code,
            "message": message
        }
    }


def get_success_response(data: any) -> dict:
    """Create standard success response."""
    return {
        "success": True,
        "data": data
    }
